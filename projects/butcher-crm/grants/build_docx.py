# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "맑은 고딕"

doc = Document()

# 기본 폰트 (한글 포함)
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(10)
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

# 여백
sec = doc.sections[0]
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)

def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        run.font.color.rgb = color

def add_para(text="", size=10, bold=False, align=None, space_after=4, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align: p.alignment = align
    if text:
        r = p.add_run(text); set_font(r, size, bold, color)
    return p

def section_bar(text):
    """■ 섹션 제목 (음영 배경)"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); set_font(r, 12, True, RGBColor(0xFF,0xFF,0xFF))
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),"404040")
    p._p.get_or_add_pPr().append(shd)
    return p

def sub_title(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text); set_font(r, 10.5, True, RGBColor(0x1F,0x3A,0x5F))
    return p

def shade_cell(cell, fill="D9D9D9"):
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"),"clear"); shd.set(qn("w:fill"),fill)
    cell._tc.get_or_add_tcPr().append(shd)

def cell_text(cell, text, bold=False, size=9.5, align=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1); p.paragraph_format.space_before = Pt(1)
    if align: p.alignment = align
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        if i>0: p.add_run().add_break()
        r = p.add_run(ln); set_font(r, size, bold, color)

def make_table(rows, cols, widths=None):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Cm(w)
    return t

# ============ 표지 제목 ============
add_para("대덕특구 기업 애로해결 바우처 지원사업", 13, True, WD_ALIGN_PARAGRAPH.CENTER, 2)
add_para("참 가 계 획 서 (붙임2)", 16, True, WD_ALIGN_PARAGRAPH.CENTER, 6)
add_para("2026년 이노폴리스캠퍼스 사업 — 원스톱 지원센터 운영사업 (1차)",
         10, False, WD_ALIGN_PARAGRAPH.CENTER, 10, RGBColor(0x60,0x60,0x60))

# 상단 기업 요약
t = make_table(2, 4, [3.0, 5.5, 3.0, 5.5])
hdr = [("기업명","주식회사 위브원","대표자","안동철"),
       ("서비스명","또와 (Ttowa)","사업자등록번호","446-81-03726")]
for ri,(a,b,c,d) in enumerate(hdr):
    cell_text(t.rows[ri].cells[0], a, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[ri].cells[0])
    cell_text(t.rows[ri].cells[1], b)
    cell_text(t.rows[ri].cells[2], c, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[ri].cells[2])
    cell_text(t.rows[ri].cells[3], d)

# ============ ■ 바우처 활용 계획 ============
section_bar("■ 바우처 활용 계획")

t = make_table(3, 2, [4.0, 13.0])
rows = [("바우처 활용기간","2026. 07. ~ 2026. 10.  (선정 통보일 ~ 2026.10, 심의 결과에 따라 조정)"),
        ("지원유형","☑ 사업화 지원형        ☐ 전문가 자문형"),
        ("희망 지원분야","① 기술·R&D 애로해결  (시제품 제작 및 PoC 지원 — 지원한도 500만원)")]
for ri,(a,b) in enumerate(rows):
    cell_text(t.rows[ri].cells[0], a, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[ri].cells[0])
    cell_text(t.rows[ri].cells[1], b)

# 과제개요
sub_title("• 과제개요 (기업 및 사업현황 / 주요 애로사항 및 발생 원인)")
add_para("[ 기업 및 사업현황 ]", 10, True, space_after=2)
add_para("㈜위브원은 정육점 특화 Zero-Task AI CRM 에이전트 「또와」를 개발·운영하는 대덕특구(대전 유성구) "
         "소재 기술창업기업입니다. 또와는 정육점 POS 거래 데이터만 연동되면 ① 단골 자동 분류·개인화 알림 발송, "
         "② 재고 분석 기반 자동 마케팅, ③ 판매 예측 기반 발주 제안을 사장님 행동 부담 ‘주 5초’로 수행하는 SaaS입니다.")
for b in [
    "정육점 매출의 70% 이상이 단골 반복구매(파일럿 매장 거래 데이터)이나, 사장님은 단골을 기억에 의존해 관리하지 못한 채 이탈·매출 손실 발생.",
    "현재 정육점 1개소 파일럿 운영 중 — 누적 거래 약 5만 건, 누적 단골 약 2,100명(5회+ 방문) 기준 단골 자동 분류·메시지 생성·알림 발송 워크플로 가동.",
    "협력사 ㈜더담우(연 매출 1,000억 정육 유통사)와 5개 점포 영업망 협력 합의, 베타 순차 도입 예정.",
    "국내 정육점 약 57,000개소(시장 약 2,200억원)를 1차 시장으로 함.",
]:
    add_para("– "+b, 9.5, space_after=2)

add_para("[ 주요 애로사항 및 애로 발생 원인 ]", 10, True, space_after=2)
t = make_table(3, 2, [3.0, 14.0])
ae = [("애로 ①","파일럿(N=1)에서 워크플로는 검증됐으나, 다수 매장(베타 N=30)으로 확장 가능한 안정화된 시제품(베타 버전)이 미완성"),
      ("애로 ②","POS 연동·알림톡 발송·효과측정(도입 전/후 매출 비교) 인프라가 PoC 수준에 머물러, 매장별 도입 전/후 효과를 정량 측정할 제품 환경 부족"),
      ("발생 원인","초기 창업기업으로 개발 인력·외주 비용·테스트 자원이 제한적이어서, 핵심 가설(매출 +15% / 비용 −50%)을 실증할 시제품 고도화 및 현장 PoC 검증 재원이 부족")]
for ri,(a,b) in enumerate(ae):
    cell_text(t.rows[ri].cells[0], a, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[ri].cells[0])
    cell_text(t.rows[ri].cells[1], b)
add_para("→ 본 바우처(시제품 제작·PoC)를 통해 베타 시제품을 완성하고 협력 매장에서 PoC를 수행하여, "
         "제품 가치 가설을 현장 데이터로 실증하는 것이 목표입니다.", 9.5, space_after=4)

# 경영정보
sub_title("• 경영정보")
t = make_table(4, 5, [3.0, 4.0, 4.0, 3.0, 3.0])
heads = ["연도","매출액(백만원)","수출액(원)","총 근로자수","신규고용"]
for i,h in enumerate(heads):
    cell_text(t.rows[0].cells[i], h, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[0].cells[i])
for ri,yr in enumerate(["2023","2024","2025"]):
    cell_text(t.rows[ri+1].cells[0], yr, align=WD_ALIGN_PARAGRAPH.CENTER)
    for ci in range(1,5):
        cell_text(t.rows[ri+1].cells[ci], "", align=WD_ALIGN_PARAGRAPH.CENTER)

# 추진역량
sub_title("• 추진역량")
add_para("보유 자료: 또와 IR 덱(17장)·사업계획서, 파일럿 매장 누적 거래 데이터(약 5만 건)·단골 데이터(약 2,100명), 회사소개서.", 9.5, space_after=2)
add_para("전담인력 / 추진체계:", 9.5, True, space_after=1)
for b in [
    "대표이사 안동철(㈜스낵포 CTO 8년, B2B SaaS 0→1 전 과정) — 과제 총괄",
    "AI·데이터 총괄 배현혜(Columbia Univ. 박사, 빅데이터 전임연구원) — 추천 엔진·효과측정",
    "제품·UX 총괄 SEO CHARLES(KAIST 학·석사) — 시제품 제작·UX",
    "운영·재무 박동일(본 사업 신청 담당) — 바우처 집행·결과보고",
    "현장·도메인 자문 이지백(정육 유통사 사외이사) — PoC 매장 연계",
]:
    add_para("– "+b, 9.5, space_after=1)
add_para("실행 준비 수준: 파일럿 1매장에서 핵심 워크플로 이미 가동 중 — 시제품 고도화 및 PoC 매장 확장만으로 즉시 착수 가능한 단계.", 9.5, space_after=4)

# 지원 필요성
sub_title("• 지원 필요성 (외부 전문가 활용 및 바우처 지원 필요 사유)")
for b in [
    "또와의 핵심 가치 가설(매출 +15% / 마케팅·재고비 −50%)은 아직 실증 데이터가 없는 N=1 단계로, 가설을 숫자로 증명할 시제품 고도화 + 다매장 PoC가 사업화의 최우선 과제.",
    "초기 창업기업 특성상 외주 개발·디자인, 알림 발송 인프라, PoC 매장 셋업 비용을 자체 부담하기에 재원이 부족.",
    "본 바우처(시제품 제작·PoC 지원, 500만원)는 이 갭을 직접 메워 ‘제품이 매장 매출을 실제로 올리는가’를 검증하는 데 결정적.",
]:
    add_para("– "+b, 9.5, space_after=2)

# 바우처 활용 계획
sub_title("• 바우처 활용 계획 — 1) 사업화 지원형")
add_para("[ 수행 내용 ]", 10, True, space_after=1)
for b in [
    "베타 시제품(또와 v1) 제작·고도화 — POS 거래 연동 모듈, 단골 자동 분류 엔진, 정육점 말투 개인화 알림톡 자동 생성·발송, 주간 효과 리포트 대시보드.",
    "효과측정 인프라 구축 — 매장별 도입 전/후 월매출 비교, 알림 1건→14일 내 재방문 전환율, 이탈군 잔존율 자동 집계 모듈.",
    "현장 PoC 수행 — 파일럿 1매장 + 협력사(㈜더담우) 연계 매장 대상 시제품 도입 및 위 지표 실측.",
]:
    add_para("– "+b, 9.5, space_after=1)
add_para("[ 제작 결과물 ]", 10, True, space_after=1)
add_para("– 또와 베타 버전 시제품(POS 연동·단골 자동관리·알림 발송·효과측정 대시보드 포함)", 9.5, space_after=1)
add_para("– PoC 검증 결과 리포트(도입 전/후 매출 비교, 재방문 전환, 잔존율 등 측정 데이터)", 9.5, space_after=1)
add_para("[ 활용 계획 ]", 10, True, space_after=1)
for b in [
    "PoC 결과를 베타 N=30 확장 및 유료 전환(PMF) 의사결정의 근거로 활용.",
    "Pre-Seed 투자유치 및 후속 정부 R&D의 핵심 실증 자료로 활용.",
    "협력사 더담우 영업망을 통한 추가 매장 도입 영업 자료로 활용.",
]:
    add_para("– "+b, 9.5, space_after=1)

# 희망 전문가
sub_title("• 희망 전문가 (컨설턴트)")
add_para("사업화 지원형 신청으로 별도 자문 전문가 미기재 — 시제품 제작·PoC 수행기관은 수행기관(대전창조경제혁신센터/"
         "이노폴리스벤처협회)에서 분야별 연계 받기를 희망. (직접 제안할 외주 개발사·디자인사가 있을 경우 성명·소속·전문분야·주요경력 기재 + 증빙)", 9.5, space_after=4)

# 제출 가능 결과물
sub_title("• 제출 가능 결과물 (사업화 지원형)")
add_para("☐ IR자료   ☐ 홍보영상   ☐ 브로셔·리플렛   ☐ 인증서   ☑ 시제품   ☑ 시험·분석 성적서(효과 측정 리포트)   ☐ 기타(        )", 10, space_after=4)

# 기대효과 및 목표
sub_title("• 기대효과 및 목표 (정량)")
t = make_table(6, 2, [4.5, 12.5])
ke = [("구분","목표 (가설 — PoC로 실증)"),
      ("제품","베타 시제품(또와 v1) 완성 + PoC 매장 도입"),
      ("매출 효과 검증","PoC 매장 도입 전/후 월매출 +10%↑ 측정"),
      ("재방문 전환","알림 1건 → 14일 내 재방문 전환율 8%↑ 측정"),
      ("사업화 연계","PoC 결과 기반 베타 30매장 확장 / 유료 전환 의사결정"),
      ("투자·고용","Pre-Seed 라운드 실증자료 확보, 6개월 내 개발·영업 인력 채용")]
for ri,(a,b) in enumerate(ke):
    bold = (ri==0)
    cell_text(t.rows[ri].cells[0], a, True, align=WD_ALIGN_PARAGRAPH.CENTER)
    cell_text(t.rows[ri].cells[1], b, bold)
    if ri==0:
        shade_cell(t.rows[ri].cells[0]); shade_cell(t.rows[ri].cells[1])

# 유사·연계 지원사업
sub_title("• 유사·연계 지원사업 참여현황")
add_para("(정부 R&D 사업) 현재 또와의 핵심 기술 개발이 진행 단계에 있으며 제품 고도화를 진행 중입니다. "
         "다만 사업화 및 시장 진입(다매장 PoC·효과 실증) 단계에서 자원 한계가 있어, 본 애로해결 바우처(시제품·PoC)를 "
         "통해 실질적인 사업화 성과(현장 실증 데이터) 창출을 도모하고자 합니다.", 9.5, space_after=2)
add_para("※ 진행 중 정부과제명·과제번호·수행기간·정부지원금은 중복지원 점검 대상이므로 사실대로 기재.", 8.5, space_after=4, color=RGBColor(0x80,0x80,0x80))

# ============ ■ 지원대상 제품·서비스 개요 ============
section_bar("■ 지원대상 제품·서비스 개요")

t = make_table(3, 2, [4.0, 13.0])
po = [("제품명","또와(Ttowa) — 정육점 특화 Zero-Task AI CRM 에이전트"),
      ("기술 개발단계","☐ 기획     ☑ 개발중     ☐ 개발완료     ☐ 사업화·상용화"),
      ("사업화 현황","☐ 투자유치 이력   ☑ PoC 진행 이력   ☐ 실증사업 참여   ☑ 고객사 확보(파일럿 1매장 + 더담우 5점포)   ☑ 정부지원사업 수행   ☐ 기타")]
for ri,(a,b) in enumerate(po):
    cell_text(t.rows[ri].cells[0], a, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[ri].cells[0])
    cell_text(t.rows[ri].cells[1], b)

sub_title("• 제품 설명")
add_para("또와는 정육점 사장님의 ‘첫 AI 직원’입니다. POS 거래만 연동되면 단골 명단·구매 이력이 자동 누적되고, "
         "AI가 이번 주 챙길 단골을 자동 선정해 정육점 말투의 개인화 메시지를 자동 작성·발송하며, 재고 분석 기반 자동 "
         "마케팅과 판매 예측 기반 발주 제안까지 수행합니다.", 9.5, space_after=2)
for b in [
    "용도/적용 분야: 동네 정육점(1~3개 매장)의 고객관리·마케팅·재고·발주 자동화. 인접 식품 소매(반찬가게·마트 등)로 확장 가능.",
    "대상 시장: 국내 정육점 약 57,000개소(약 2,200억원) → 식품 소매·외식(약 90만 개) → 재방문 기반 B2C 소상공인 전반.",
    "주요 경쟁력·차별성: ① 정육 수직 도메인 완전 특화(부분육·계절·명절·가구단위 학습 데이터), ② Zero-Task UX(사장님 손=0, 주 5초), ③ POS 교체 불필요(사장님 직접 가입). 토스플레이스·캐시노트 등 수평 SaaS가 진입하지 못한 정육 수직 카테고리를 선점.",
]:
    add_para("– "+b, 9.5, space_after=2)

sub_title("• 관련 제품·서비스 이미지")
add_para("(또와 앱/대시보드 스크린샷·컨셉 이미지 첨부 위치)", 9.5, space_after=4, color=RGBColor(0x80,0x80,0x80))

sub_title("• 관련 지식재산권")
t = make_table(6, 2, [4.5, 12.5])
ip = [("특허(출원)","                                              (출원번호:                    )"),
      ("특허(등록)","                                              (등록번호:                    )"),
      ("인증(GS/KC/기타)",""),
      ("상표권 (‘또와’)","                                              (번호:                    )"),
      ("프로그램 등록(SW)","                                              (번호:                    )"),
      ("기타","")]
for ri,(a,b) in enumerate(ip):
    cell_text(t.rows[ri].cells[0], a, True); shade_cell(t.rows[ri].cells[0])
    cell_text(t.rows[ri].cells[1], b)

sub_title("• 최근 3개년(’23~’25) 주요 판매실적")
t = make_table(4, 3, [5.0, 4.0, 8.0])
for i,h in enumerate(["판매처","판매금액(백만원)","주요내용 (판매일자/거래내용)"]):
    cell_text(t.rows[0].cells[i], h, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[0].cells[i])
for ri in range(1,4):
    for ci in range(3):
        cell_text(t.rows[ri].cells[ci], "")
add_para("※ 기재 내용 관련 증빙자료(특허·인증·투자·계약·거래내역 등)는 PDF로 별도 첨부.", 8.5, space_after=6, color=RGBColor(0x80,0x80,0x80))

# ============ 붙임1 참여신청서 (페이지 분리) ============
doc.add_page_break()
add_para("[참고] 붙임1. 참여신청서 — 기재안", 13, True, WD_ALIGN_PARAGRAPH.CENTER, 6)
add_para("※ 공식 붙임1 양식에 아래 내용을 옮겨 기재 후 대표자 서명. 빈칸은 확인 후 기재.",
         9, False, WD_ALIGN_PARAGRAPH.CENTER, 8, RGBColor(0x60,0x60,0x60))

t = make_table(16, 2, [5.0, 12.0])
ap = [
    ("기업명","주식회사 위브원"),
    ("대표자명","안동철"),
    ("사업자등록번호","446-81-03726"),
    ("법인설립일자",""),
    ("대표자 연락처(HP)","010-4301-8469"),
    ("대표자 이메일","ahn@weaveone.kr"),
    ("신청자명/직위","박동일 / 팀장"),
    ("신청자 연락처","(이메일) park@weaveone.kr  /  (C.P.) 010-4301-8469"),
    ("본사주소","대전광역시 유성구 대학로 157, 302호  (우편번호:           )"),
    ("홈페이지",""),
    ("주요제품","정육점 특화 AI CRM SaaS 「또와」"),
    ("사업자 구분","☐ 개인     ☑ 법인"),
    ("업종 / 업태",""),
    ("지원유형","☑ 사업화 지원형     ☐ 전문가 자문형"),
    ("희망 지원분야","☑ 기술·R&D 애로해결  (1개 분야 선택)"),
    ("희망 지원내용","또와 베타 시제품 제작 및 정육 매장 현장 PoC 검증 (효과측정 인프라 포함)"),
]
for ri,(a,b) in enumerate(ap):
    cell_text(t.rows[ri].cells[0], a, True, align=WD_ALIGN_PARAGRAPH.CENTER); shade_cell(t.rows[ri].cells[0])
    cell_text(t.rows[ri].cells[1], b)

add_para("", space_after=4)
add_para("사업안내 기관:   ☐ 대전창조경제혁신센터     ☐ 이노폴리스벤처협회   (접수 경로에 맞게 택1)", 9.5, space_after=8)
add_para("당사는 상기와 같이 원스톱 지원센터 운영사업 애로해결 바우처 참가를 신청하며, 지원대상으로 선정 시 "
         "본 사업과 관련한 업무수행 및 제반 자료제출 등을 성실히 수행할 것을 확약합니다.", 9.5, space_after=10)
add_para("2026년        월        일", 10, True, WD_ALIGN_PARAGRAPH.CENTER, 6)
add_para("대표자(신청인) :  안 동 철        (서명 또는 인)", 10, False, WD_ALIGN_PARAGRAPH.RIGHT, 4)

out = "/home/user/Park_Home/projects/butcher-crm/grants/대덕특구_애로해결바우처_참가계획서_v0.1.docx"
doc.save(out)
print("saved:", out)
