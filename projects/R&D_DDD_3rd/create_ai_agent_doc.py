from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# 기본 스타일 설정
style = doc.styles['Normal']
style.font.name = '맑은 고딕'
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

# 여백 설정
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = '맑은 고딕'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    return heading

def add_para(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '맑은 고딕'
    run.font.size = Pt(11)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    return p

def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.name = '맑은 고딕'
                run.font.size = Pt(10)
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    for row_data in rows:
        row_cells = table.add_row().cells
        for i, cell_text in enumerate(row_data):
            row_cells[i].text = cell_text
            for paragraph in row_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.name = '맑은 고딕'
                    run.font.size = Pt(10)
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    return table

# ========== 문서 작성 ==========

# 제목
title = doc.add_heading('고객관리 AI 에이전트', level=0)
for run in title.runs:
    run.font.name = '맑은 고딕'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

subtitle = doc.add_heading('기능 및 기술 스택 정의서', level=1)
for run in subtitle.runs:
    run.font.name = '맑은 고딕'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')

doc.add_paragraph()

# ========== 1. 핵심 기능 영역 ==========
add_heading(doc, '1. 핵심 기능 영역', level=1)

# 1.1
add_heading(doc, '1.1 고객 데이터 통합 (Customer Data Platform)', level=2)
add_table(doc, ['기능', '설명'], [
    ['멀티채널 데이터 수집', 'POS, 웹, 앱, 전화, SNS 등 접점별 데이터 통합'],
    ['고객 ID 통합 (Identity Resolution)', '여러 채널의 동일 고객을 하나로 매핑'],
    ['실시간 이벤트 스트리밍', '구매/방문/클릭 등 행동 데이터 실시간 수집'],
    ['데이터 정제/표준화', '결측치 처리, 이상치 탐지, 스키마 통일'],
])
doc.add_paragraph()
add_para(doc, '기술 스택: Apache Airflow, dbt, Fivetran (ETL/ELT) | Apache Kafka, AWS Kinesis (스트리밍) | Snowflake, BigQuery, PostgreSQL (스토리지) | Segment, mParticle (CDP)')

# 1.2
add_heading(doc, '1.2 고객 분석 및 세분화 (Segmentation & Profiling)', level=2)
add_table(doc, ['기능', '설명'], [
    ['RFM 분석', 'Recency, Frequency, Monetary 기반 고객 등급화'],
    ['고객 클러스터링', '행동 패턴 기반 자동 세그먼트 생성'],
    ['고객 프로파일링', '인구통계 + 행동 + 선호도 통합 프로필'],
    ['코호트 분석', '가입/첫 구매 시점별 집단 행동 추적'],
])
doc.add_paragraph()
add_para(doc, '기술 스택/방법론: K-Means, DBSCAN, Hierarchical Clustering (클러스터링) | PCA, t-SNE, UMAP (차원 축소/시각화) | scikit-learn, pandas')

# 1.3
add_heading(doc, '1.3 예측 모델링 (Predictive Analytics)', level=2)
add_table(doc, ['기능', '설명'], [
    ['이탈 예측 (Churn Prediction)', '이탈 확률 스코어링, 이탈 시점 예측'],
    ['생애가치 예측 (CLV/LTV)', '고객별 미래 기대 수익 산출'],
    ['구매 확률 예측', '다음 구매 시점, 구매 상품 예측'],
    ['반응 예측', '마케팅 메시지에 대한 반응 확률'],
])
doc.add_paragraph()
add_para(doc, '기술 스택/방법론: Random Forest, XGBoost, LightGBM, CatBoost (분류) | Cox Proportional Hazard, Kaplan-Meier (생존 분석) | Prophet, ARIMA (시계열) | LSTM, Transformer (딥러닝 시퀀스) | H2O, AutoGluon (AutoML)')

# 1.4
add_heading(doc, '1.4 개인화 커뮤니케이션 (Personalized Outreach)', level=2)
add_table(doc, ['기능', '설명'], [
    ['개인화 메시지 생성', '고객별 맞춤 문자/이메일/푸시 자동 작성'],
    ['최적 발송 시점 결정', '고객별 반응률 높은 시간대 예측'],
    ['채널 최적화', '문자 vs 카카오 vs 이메일 vs 앱푸시 선택'],
    ['A/B 테스트 자동화', '메시지 변형 자동 생성 및 성과 비교'],
])
doc.add_paragraph()
add_para(doc, '기술 스택/방법론: GPT-4, Claude, Gemini, LLaMA, Mistral (LLM) | Few-shot, Chain-of-Thought, RAG (프롬프트 엔지니어링) | Multi-Armed Bandit, Contextual Bandit (발송 최적화) | Reinforcement Learning (채널 선택)')

# 1.5
add_heading(doc, '1.5 추천 시스템 (Recommendation)', level=2)
add_table(doc, ['기능', '설명'], [
    ['상품 추천', '구매 이력 기반 다음 구매 상품 추천'],
    ['번들 추천', '함께 구매하면 좋은 상품 조합'],
    ['컨텐츠 추천', '관심사 기반 블로그/영상 추천'],
    ['Next Best Action', '지금 이 고객에게 가장 효과적인 액션 제안'],
])
doc.add_paragraph()
add_para(doc, '기술 스택/방법론: Matrix Factorization, ALS (협업 필터링) | TF-IDF, Embedding Similarity (콘텐츠 기반) | Neural Collaborative Filtering, Two-Tower Model (딥러닝) | GNN (그래프 기반) | Surprise, LensKit, TensorFlow Recommenders')

# 1.6
add_heading(doc, '1.6 대화형 인터페이스 (Conversational AI)', level=2)
add_table(doc, ['기능', '설명'], [
    ['고객 문의 자동 응대', 'FAQ, 주문 조회, 예약 등 자동 처리'],
    ['자연어 질의 분석', '고객 요청 의도 파악 (Intent Classification)'],
    ['멀티턴 대화 관리', '맥락 유지하며 여러 턴 대화'],
    ['감정 분석', '고객 불만/긍정 감지 및 에스컬레이션'],
])
doc.add_paragraph()
add_para(doc, '기술 스택/방법론: LangChain, LlamaIndex (LLM + RAG) | spaCy, Hugging Face Transformers (Intent/NER) | Rasa, Dialogflow (대화 관리) | BERT 기반 Sentiment Classification (감정 분석) | Whisper (STT), TTS API (음성)')

# 1.7
add_heading(doc, '1.7 자동화 및 워크플로우 (Automation & Orchestration)', level=2)
add_table(doc, ['기능', '설명'], [
    ['트리거 기반 자동화', '특정 조건 충족 시 자동 액션 실행'],
    ['캠페인 오케스트레이션', '멀티스텝 마케팅 캠페인 자동 운영'],
    ['재학습 자동화', '모델 성능 모니터링 및 자동 재학습'],
    ['알림/에스컬레이션', '이상 징후 감지 시 담당자 알림'],
])
doc.add_paragraph()
add_para(doc, '기술 스택: Apache Airflow, Prefect, Temporal (워크플로우) | MLflow, Kubeflow, Weights & Biases (MLOps) | Grafana, Prometheus, Evidently AI (모니터링/드리프트) | Slack/Webhook (알림)')

# 1.8
add_heading(doc, '1.8 분석 대시보드 (Analytics & Insights)', level=2)
add_table(doc, ['기능', '설명'], [
    ['고객 현황 대시보드', '세그먼트별 현황, 이탈 위험 고객 리스트'],
    ['캠페인 성과 분석', '발송/오픈/클릭/전환/매출 기여도'],
    ['인사이트 자동 생성', 'LLM 기반 주간 리포트 자동 작성'],
    ['셀프서비스 분석', '비개발자도 드래그앤드롭으로 분석'],
])
doc.add_paragraph()
add_para(doc, '기술 스택: Metabase, Superset, Tableau, Looker (BI 도구) | Streamlit, Gradio (임베디드 분석) | Text-to-SQL with LangChain + SQLAgent (자연어 질의)')

doc.add_page_break()

# ========== 2. 전체 아키텍처 ==========
add_heading(doc, '2. 전체 아키텍처 개요', level=1)

add_para(doc, '고객관리 AI 에이전트의 전체 시스템 아키텍처는 다음 5개 레이어로 구성된다.')
doc.add_paragraph()

add_table(doc, ['레이어', '구성 요소', '역할'], [
    ['데이터 소스', 'POS, 웹/앱, CRM, SNS, 콜센터, 외부 데이터', '원천 데이터 발생 지점'],
    ['데이터 파이프라인', 'ETL/스트리밍 → 정제 → 통합 → Feature Store', '데이터 수집, 변환, 저장'],
    ['AI/ML 엔진', '세분화, 이탈예측, CLV, 추천, NLP/LLM, 최적화', '분석 및 예측 모델 실행'],
    ['액션 레이어', '메시지 발송, 챗봇 응대, 캠페인 실행, 알림', '고객 접점 자동화 실행'],
    ['사용자 인터페이스', '사장님 대시보드, 관리자 콘솔, 3rd Party API', '사용자/시스템 상호작용'],
])

doc.add_paragraph()
add_para(doc, '각 레이어는 API 기반으로 느슨하게 결합(Loosely Coupled)되어 개별 모듈의 교체·확장이 용이하도록 설계한다. 특히 AI/ML 엔진은 모델별로 마이크로서비스 형태로 분리하여 독립적인 배포·스케일링이 가능하도록 구성한다.')

doc.add_page_break()

# ========== 3. 핵심 방법론 ==========
add_heading(doc, '3. 핵심 방법론 요약', level=1)

add_table(doc, ['영역', '방법론', '활용 목적'], [
    ['고객 가치', 'RFM 분석, CLV 모델링', '고객 등급화, 투자 우선순위 결정'],
    ['이탈 예측', 'Survival Analysis, Gradient Boosting', '이탈 확률 및 시점 예측'],
    ['개인화', 'Collaborative Filtering, LLM Prompt Engineering', '상품 추천, 메시지 생성'],
    ['최적화', 'Multi-Armed Bandit, Reinforcement Learning', '발송 시점, 채널, 메시지 최적화'],
    ['검증', 'A/B Test, RCT, Causal Inference', '마케팅 효과 인과적 검증'],
    ['자동화', 'MLOps, CI/CD for ML', '모델 배포, 재학습 자동화'],
])

doc.add_paragraph()

add_heading(doc, '3.1 RFM 분석', level=2)
add_para(doc, 'Recency(최근성), Frequency(빈도), Monetary(금액)의 세 축으로 고객을 평가하는 전통적이지만 강력한 세분화 기법이다. 각 축을 5분위로 나누어 125개 세그먼트를 생성하거나, 점수를 합산하여 VIP/일반/이탈위험 등으로 단순화할 수 있다.')

add_heading(doc, '3.2 생존 분석 (Survival Analysis)', level=2)
add_para(doc, '고객이 "언제" 이탈할 것인지를 예측하는 방법론이다. Cox Proportional Hazard 모델은 각 고객 특성이 이탈 위험에 미치는 영향을 정량화하며, Kaplan-Meier 곡선은 시간에 따른 생존(잔존) 확률을 시각화한다. 단순 이탈 확률(Yes/No)을 넘어 이탈 시점 예측이 가능하여 선제적 개입 타이밍을 결정하는 데 유용하다.')

add_heading(doc, '3.3 Multi-Armed Bandit', level=2)
add_para(doc, 'A/B 테스트의 한계(고정 비율 배분으로 인한 기회비용)를 극복하는 적응형 실험 기법이다. 실시간으로 성과가 좋은 변형에 더 많은 트래픽을 배분하여 탐색(Exploration)과 활용(Exploitation)을 균형 있게 수행한다. Thompson Sampling, UCB(Upper Confidence Bound) 등의 알고리즘이 대표적이다.')

add_heading(doc, '3.4 RAG (Retrieval-Augmented Generation)', level=2)
add_para(doc, 'LLM의 환각(Hallucination) 문제를 완화하고 최신/도메인 특화 정보를 반영하기 위해, 외부 지식 베이스에서 관련 문서를 검색(Retrieval)한 후 이를 컨텍스트로 제공하여 생성(Generation)하는 기법이다. 정육점 도메인 지식(부위-용도 매핑, 계절-고기 페어링 등)을 벡터 DB에 저장하고 프롬프트에 주입하는 방식으로 활용할 수 있다.')

doc.add_page_break()

# ========== 4. 기술 스택 총괄 ==========
add_heading(doc, '4. 기술 스택 총괄표', level=1)

add_table(doc, ['카테고리', '기술/도구', '용도'], [
    ['데이터 파이프라인', 'Airflow, dbt, Kafka, Kinesis', 'ETL, 스트리밍, 배치 처리'],
    ['데이터 저장', 'Snowflake, BigQuery, PostgreSQL, Redis', 'DW, 캐시, Feature Store'],
    ['ML 프레임워크', 'scikit-learn, XGBoost, LightGBM, PyTorch', '모델 학습'],
    ['LLM/NLP', 'OpenAI API, Claude API, LangChain, Hugging Face', '텍스트 생성, 임베딩'],
    ['추천', 'TensorFlow Recommenders, Surprise, LensKit', '협업 필터링, 콘텐츠 기반'],
    ['MLOps', 'MLflow, Kubeflow, W&B, Evidently', '실험 추적, 배포, 모니터링'],
    ['백엔드', 'FastAPI, Django, Node.js', 'API 서버'],
    ['프론트엔드', 'React, Next.js, Streamlit', '대시보드, 관리 콘솔'],
    ['인프라', 'AWS, GCP, Kubernetes, Docker', '클라우드, 컨테이너 오케스트레이션'],
    ['BI/시각화', 'Metabase, Superset, Grafana', '대시보드, 모니터링'],
])

doc.add_page_break()

# ========== 5. 확장 로드맵 ==========
add_heading(doc, '5. 현재 대비 확장 가능 영역', level=1)

add_para(doc, '현재 정육점 특화 AI 고객관리 에이전트 프로젝트에서 구현 중인 기능과, 향후 확장 가능한 고도화 영역을 비교하면 다음과 같다.')

doc.add_paragraph()

add_table(doc, ['현재 구현', '향후 확장 가능'], [
    ['RF 기반 이탈 예측 (Yes/No)', 'Survival Analysis로 이탈 "시점" 예측 추가'],
    ['LLM 문자 생성 (단방향)', '챗봇/음성 에이전트로 양방향 대화 확장'],
    ['문자 단일 채널', '카카오톡, 앱푸시, 이메일 멀티채널'],
    ['수동 발송 시점 설정', 'Bandit 알고리즘으로 최적 시점 자동 결정'],
    ['월 1회 배치 재학습', '실시간/일간 재학습 + 드리프트 모니터링'],
    ['정적 대시보드', 'Text-to-SQL 자연어 질의 기능'],
    ['정육점 단일 산업', '반찬가게, 떡집, 꽃집 등 멀티 버티컬'],
])

doc.add_paragraph()
add_para(doc, '확장은 단계적으로 진행하되, 정육점에서의 PMF(Product-Market Fit) 확립 후 기능 고도화와 산업 확장을 병행하는 전략을 권장한다.')

# ========== 저장 ==========
output_path = '/home/user/Park_Home/projects/R&D_DDD_3rd/고객관리_AI_에이전트_기능_및_기술스택.docx'
doc.save(output_path)
print(f'파일 저장 완료: {output_path}')
